# -*- coding: utf-8 -*-
import base64
import io
import logging
import re
from datetime import date, datetime, timedelta

import requests

from odoo import _, api, fields, models
from odoo.exceptions import UserError
from odoo.tools.binary import BinaryBytes

_logger = logging.getLogger(__name__)

REFERENCE_COLUMNS = [
    ('vendor', 'Vendor'),
    ('product_template', 'Product Template'),
    ('vendor_product_code', 'Vendor Product Code'),
    ('date_start', 'Start Date'),
    ('date_end', 'End Date'),
    ('company', 'Company'),
    ('min_qty', 'Quantity'),
    ('unit', 'Unit'),
    ('price', 'Unit Price'),
    ('discount', 'Discount (%)'),
    ('delay', 'Lead Time'),
]
COLUMN_ALIASES = {key: [label.lower()] for key, label in REFERENCE_COLUMNS}
COLUMN_ALIASES['vendor'].append('partner_id')
COLUMN_ALIASES['vendor_product_code'].append('product_code')
COLUMN_ALIASES['product_template'] += ['product_tmpl_id', 'product_template']
COLUMN_ALIASES['date_start'].append('date_start')
COLUMN_ALIASES['date_end'].append('date_end')
COLUMN_ALIASES['company'].append('company_id')
COLUMN_ALIASES['min_qty'].append('min_qty')
COLUMN_ALIASES['price'].append('price')
COLUMN_ALIASES['discount'].append('discount')
COLUMN_ALIASES['delay'].append('delay')
COLUMN_ALIASES['note'] = ['note']

PRODUCT_CODE_RE = re.compile(r'^\[(?P<code>[^\]]+)\]\s*(?P<name>.*)$')


class VendorPricelistUpdate(models.Model):
    _name = 'vendor.pricelist.update'
    _description = 'Vendor Pricelist AI Update'
    _inherit = ['mail.thread', 'mail.activity.mixin']
    _order = 'id desc'
    _rec_name = 'reference_id'

    reference_id = fields.Char(string='Reference', default='New', readonly=True, copy=False)
    name = fields.Char(string="Name", copy=False)

    vendor_id = fields.Many2one(
        'res.partner', string='Vendor', required=True, tracking=True,
        domain=[('supplier_rank', '>', 0)], )

    company_id = fields.Many2one(
        'res.company', string='Company', default=lambda self: self.env.company,
        required=True,
    )
    state = fields.Selection([
        ('draft', 'Draft'),
        ('file_generated', 'Pricelist File Generated'),
        ('ai_done', 'AI Analysis Done'),
        ('error', 'AI Error'),
        ('applied', 'Applied'),
    ], string='Status', default='draft', tracking=True, copy=False)

    # -- Step 1: generated reference file (current Vendor Pricelist) -------
    pricelist_file = fields.Binary(string='Vendor Pricelist File', attachment=True, copy=False)
    pricelist_filename = fields.Char(string='Pricelist Filename', copy=False)
    pricelist_line_count = fields.Integer(string='Pricelist Lines', readonly=True, copy=False)
    generated_date = fields.Datetime(string='Generated On', readonly=True, copy=False)
    # -- Step 2: file uploaded back by the vendor ---------------------------
    vendor_file = fields.Binary(string="Vendor Pricelist File", attachment=True, copy=False)
    vendor_filename = fields.Char(string='Vendor File Name', copy=False)

    prompt = fields.Text(
        string='AI Prompt', copy=False,
        help='Prompt sent to the AI for this record. Pre-filled from the default '
             'AI Prompt (Pricelist Management > AI Prompt) when the record is '
             'created - editing it here only changes this record, it never '
             'updates the default AI Prompt.',
    )

    # -- Step 3: AI analysis --------------------------------------------------
    ai_run_date = fields.Datetime(string='AI Run Date', readonly=True, copy=False)

    error_message = fields.Text(string='Error Details', readonly=True, copy=False)

    line_ids = fields.One2many(
        'vendor.pricelist.update.line', 'update_id', string='Product Matching',
        domain=[('matched', '=', True)],
    )
    unmatched_line_ids = fields.One2many(
        'vendor.pricelist.update.line', 'update_id', string='Unmatched Products',
        domain=[('matched', '=', False)],
    )
    matched_count = fields.Integer(compute='_compute_line_counts', string='Matched')
    unmatched_count = fields.Integer(compute='_compute_line_counts', string='Unmatched')

    applied_date = fields.Datetime(string='Applied On', readonly=True, copy=False)

    ai_response = fields.Text(
        string='AI Response', readonly=True, copy=False,
        help='Text message returned by the AI along with the generated file.',
    )

    ai_result_file = fields.Binary(string='AI Result File', attachment=True, copy=False)
    ai_result_filename = fields.Char(string='AI Result Filename', copy=False)
    ai_total_tokens = fields.Integer(string='AI Total Tokens Used', readonly=True, copy=False)

    @api.depends('line_ids.matched', 'unmatched_line_ids.matched')
    def _compute_line_counts(self):
        for rec in self:
            rec.matched_count = len(rec.line_ids)
            rec.unmatched_count = len(rec.unmatched_line_ids)

    @api.model
    def default_get(self, fields_list):
        res = super().default_get(fields_list)
        if 'prompt' in fields_list and not res.get('prompt'):
            default_prompt = self.env['ai.prompt.vts'].get_prompt_text()
            if default_prompt:
                res['prompt'] = default_prompt.replace('{{COMPANY_NAME}}', self.env.company.name or '')
        return res

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            if not vals.get('reference_id') or vals['reference_id'] == 'New':
                vals['reference_id'] = self.env['ir.sequence'].next_by_code('vendor.pricelist.update') or 'New'
        return super().create(vals_list)

    def unlink(self):
        Attachment = self.env['ir.attachment'].sudo()
        attachments = Attachment.search([('res_model', '=', self._name),('res_id', 'in', self.ids),])
        attachments.unlink()
        return super().unlink()
    # =========================================================================
    # STEP 1 - Generate the current Vendor Pricelist file for the vendor
    # =========================================================================
    def action_generate_pricelist_file(self):
        self.ensure_one()
        if not self.vendor_id:
            raise UserError(_('Please select a Vendor first.'))

        supplierinfos = self.env['product.supplierinfo'].sudo().search([
            ('partner_id', '=', self.vendor_id.id),
        ])

        xlsx_data = self._build_pricelist_xlsx(supplierinfos)
        filename = 'SUPPLIER PRICELIST IMPORT EXPORT.xlsx'

        self.write({
            'pricelist_file': BinaryBytes(xlsx_data, filename=filename),
            'pricelist_filename': filename,
            'pricelist_line_count': len(supplierinfos),
            'generated_date': fields.Datetime.now(),
            'state': 'file_generated',
        })
        self.message_post(
            body=_('Vendor Pricelist file generated with %s line(s) for %s.') % (
                len(supplierinfos), self.vendor_id.display_name),
        )
        return True

    def _build_pricelist_xlsx(self, supplierinfos):
        """Build the XLSX content, columns in the order/labels of
        REFERENCE_COLUMNS."""
        try:
            import openpyxl
            from openpyxl.styles import Font
        except ImportError:
            raise UserError(_('The "openpyxl" python library is required to generate the pricelist file.'))

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Vendor Pricelist'

        headers = [label for _key, label in REFERENCE_COLUMNS]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)

        for line in supplierinfos:
            row = [
                line.partner_id.display_name or '',
                line.product_tmpl_id.display_name or '',
                line.product_code or '',
                line.date_start.strftime('%Y-%m-%d') if line.date_start else '',
                line.date_end.strftime('%Y-%m-%d') if line.date_end else '',
                line.company_id.display_name if line.company_id else '',
                line.min_qty,
                line.uom_id.display_name if line.uom_id else '',
                line.price,
                line.discount if hasattr(line, 'discount') else 0.0,
                line.delay,
            ]
            ws.append(row)

        for col_idx in range(1, len(headers) + 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 22

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    # =========================================================================
    # STEP 3 - AI Analysis
    # =========================================================================

    @staticmethod
    def _extract_output_file_id(data):
        for item in data.get('output', []):
            if item.get('type') != 'message':
                continue
            for block in item.get('content', []):
                if block.get('type') == 'output_text':
                    for ann in block.get('annotations', []):
                        if ann.get('type') == 'container_file_citation':
                            return ann.get('file_id'), ann.get('container_id'), ann.get('filename')
        return None, None, None

    @staticmethod
    def _extract_output_text(data):
        texts = []
        for item in data.get('output', []):
            if item.get('type') != 'message':
                continue
            for block in item.get('content', []):
                if block.get('type') == 'output_text' and block.get('text'):
                    texts.append(block['text'])
        return '\n'.join(texts)

    def action_run_ai_analysis(self):
        self.ensure_one()
        if not self.pricelist_file:
            raise UserError(_('Please generate the Vendor Pricelist file first.'))
        if not self.vendor_file:
            raise UserError(_("Please upload the vendor's returned price file first."))

        config = self.env['vraja.ai.config'].sudo().search([], limit=1)
        if not config:
            raise UserError(_('Vraja AI configuration not found. Please configure it in AI Dashboard > Configuration.'))
        if config.ai_provider != 'openai':
            raise UserError(_('Please select OpenAI as the AI Provider in the Vraja AI configuration.'))
        openai_api_key = (config.openai_api_key or '').strip()
        if not openai_api_key:
            raise UserError(_('No OpenAI API Key configured. Set it in AI Dashboard > Configuration.'))
        llm_model = config.llm_model or 'gpt-5.4'

        # Clear any previous run.
        (self.line_ids | self.unmatched_line_ids).unlink()
        self.write({'error_message': False, 'ai_response': False})

        try:
            vendor_raw_bytes = self.vendor_file.content
            vendor_file_id = self._upload_file_to_openai(openai_api_key, vendor_raw_bytes, self.vendor_filename)

            pricelist_raw_bytes = self.pricelist_file.content
            pricelist_file_id = self._upload_file_to_openai(openai_api_key, pricelist_raw_bytes, self.pricelist_filename)

            if not self.prompt:
                raise UserError(_('No AI Prompt set on this record. Please enter the required prompt instruction for AI'))
            prompt = self.prompt.replace('{{COMPANY_NAME}}', self.company_id.name or '')

            result_file_id, container_id, ai_filename, ai_usage, ai_text = self._call_openai(
                openai_api_key, llm_model, prompt, [pricelist_file_id, vendor_file_id])

            result_bytes = self.download_openai_file(openai_api_key, result_file_id, container_id)

            result_filename = ai_filename or (
                    'AI_Import_Result_%s.xlsx' % fields.Date.context_today(self).strftime('%Y%m%d'))

            self.write({
                'ai_run_date': fields.Datetime.now(),
                'ai_response': ai_text or False,
                'ai_result_file': BinaryBytes(result_bytes, filename=result_filename),
                'ai_result_filename': result_filename,
                'ai_total_tokens': ai_usage.get('total_tokens') or 0,
            })

            self._create_lines_from_ai(result_bytes)

            self.write({
                'state': 'ai_done',
                'error_message': False,
            })

            attachment = self.env['ir.attachment'].create({
                'name': result_filename,
                'raw': result_bytes,
                'res_model': self._name,
                'res_id': self.id,
                'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            })
            self.message_post(
                body=_(
                    'AI analysis completed: %s matched, %s unmatched product line(s).'
                ) % (self.matched_count, self.unmatched_count),
                attachment_ids=[attachment.id],
            )

        except Exception as exc:
            _logger.exception(_('Vendor Pricelist AI Update: AI analysis failed'))
            self.write({
                'state': 'error',
                'error_message': str(exc),
            })
            self.message_post(body=_('AI analysis failed: %s') % str(exc))

        return True

    def _upload_file_to_openai(self, openai_api_key, raw_bytes, filename):
        resp = requests.post(
            'https://api.openai.com/v1/files',
            headers={'Authorization': 'Bearer %s' % openai_api_key},
            files={'file': (filename, raw_bytes)},
            data={'purpose': 'user_data'},
            timeout=60,
        )
        self._raise_for_http_error(resp, _('OpenAI (file upload)'))
        return resp.json()['id']

    def _call_openai(self, openai_api_key, llm_model, prompt, file_id=None):
        if not openai_api_key:
            raise UserError(_('OpenAI API Key is not configured.'))

        content = [{'type': 'input_text', 'text': prompt}]
        file_ids = file_id if isinstance(file_id, (list, tuple)) else ([file_id] if file_id else [])
        for fid in file_ids:
            content.append({'type': 'input_file', 'file_id': fid})

        resp = requests.post(
            'https://api.openai.com/v1/responses',
            headers={
                'Authorization': 'Bearer %s' % openai_api_key,
                'Content-Type': 'application/json',
            },
            json={
                'model': llm_model,
                'input': [{'role': 'user', 'content': content}],
                'tools': [{'type': 'code_interpreter', 'container': {'type': 'auto'}}],
            },
            timeout=800,
        )
       
        self._raise_for_http_error(resp, _('OpenAI'))
        data = resp.json()

        output_file_id, container_id, output_filename = self._extract_output_file_id(data)
        usage = data.get('usage') or {}

        message_text = self._extract_output_text(data)

        if not output_file_id:
            self.write({'ai_response': message_text or False})
            raise UserError(_('The AI did not return a generated file. Response text: %s') % message_text)

        return output_file_id, container_id, output_filename, usage, message_text

    def download_openai_file(self, openai_api_key, file_id, container_id=None):
        if container_id:
            url = 'https://api.openai.com/v1/containers/%s/files/%s/content' % (container_id, file_id)
        else:
            url = 'https://api.openai.com/v1/files/%s/content' % file_id
        resp = requests.get(
            url,
            headers={'Authorization': 'Bearer %s' % openai_api_key},
            timeout=60,
        )
       
        self._raise_for_http_error(resp, _('OpenAI (file download)'))
        return resp.content

    @staticmethod
    def _raise_for_http_error(resp, provider_label):
        if resp.status_code >= 400:
            raise UserError(_('%s API error (HTTP %s): %s') % (
                provider_label, resp.status_code, resp.text[:1000]))

    def _resolve_uom(self, unit_text, product_tmpl=None):
        UoM = self.env['uom.uom']
        text = (str(unit_text).strip() if unit_text else '')
        if text:
            uom = UoM.search([('name', '=ilike', text)], limit=1)
            if not uom:
                uom = UoM.search([('name', 'ilike', text)], limit=1)
            if uom:
                return uom
        if product_tmpl:
            if product_tmpl.uom_id:
                return product_tmpl.uom_id
        return (
                self.env.ref('uom.product_uom_unit', raise_if_not_found=False)
                or UoM.search([], limit=1)
        )

    @staticmethod
    def _to_float(value, default=0.0):
        if value in (None, ''):
            return default
        try:
            return float(value)
        except (TypeError, ValueError):
            return default

    @staticmethod
    def _to_date(value):
        if not value:
            return False
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        try:
            return fields.Date.to_date(str(value).strip())
        except Exception:
            return False

    @staticmethod
    def _find_sheet(wb, name):
        """Find a worksheet by name, case/whitespace-insensitive."""
        target_name = next(
            (s for s in wb.sheetnames if s.strip().lower() == name), None)
        return wb[target_name] if target_name else None

    @staticmethod
    def _sheet_rows_and_getter(ws):
        """Read a worksheet's rows and return (data_rows, get(row, key, default))
        using COLUMN_ALIASES to locate each logical column by its header,
        regardless of column order."""
        rows = list(ws.iter_rows(values_only=True))
        header = [str(h).strip().lower() if h else '' for h in rows[0]] if rows else []
        col = {}
        for key, aliases in COLUMN_ALIASES.items():
            for alias in aliases:
                if alias in header:
                    col[key] = header.index(alias)
                    break

        def get(row, key, default=None):
            idx = col.get(key)
            return row[idx] if idx is not None and idx < len(row) else default

        return rows[1:] if rows else [], get

    def _resolve_product_template(self, raw_value):
        """product_tmpl_id may come back three ways from the AI, for a
        certain match on the IMPORT ODOO sheet:
        - a raw Odoo database id (int/float/numeric string) - the common
          case, since the AI reads real ids from the source file;
        - '[default_code] Name' text;
        - a plain existing product name.
        Try each in order of reliability. Always returns a tuple
        (product_tmpl_recordset_or_empty, display_text) so the readable
        name is kept even when no Odoo product was found."""
        Template = self.env['product.template']
        if raw_value in (None, ''):
            return Template, ''

        # Raw database id.
        if isinstance(raw_value, (int, float)) and not isinstance(raw_value, bool):
            product_tmpl = Template.browse(int(raw_value)).exists()
            if product_tmpl:
                return product_tmpl, product_tmpl.name

        text = str(raw_value).strip()
        if text.isdigit():
            product_tmpl = Template.browse(int(text)).exists()
            if product_tmpl:
                return product_tmpl, product_tmpl.name

        match = PRODUCT_CODE_RE.match(text)
        default_code, name_part = (match.group('code').strip(), match.group('name').strip()) \
            if match else (None, text)
        display_text = name_part or text
        product_tmpl = Template
        if default_code:
            product_tmpl = Template.search([('default_code', '=', default_code)], limit=1)
        if not product_tmpl and name_part:
            product_tmpl = Template.search([('name', '=', name_part)], limit=1)
        if not product_tmpl:
            product_tmpl = Template.search([('name', '=', text)], limit=1)
        return product_tmpl, (product_tmpl.name if product_tmpl else display_text)

    def _create_lines_from_ai(self, xlsx_bytes: object) -> None:
        """Read the AI result file and create one vendor.pricelist.update.line
        per row, from two separate, mutually exclusive sheets:
        - "IMPORT ODOO": certain matches only -> Product Matching tab (matched=True).
        - "UNMATCHED PRODUCTS": probable/uncertain matches and genuinely new
          supplier products -> Unmatched Products tab (matched=False), with the
          exact same field completeness as matched lines so they can be mapped
          to an existing product or turned into a new product afterward.
        """
        self.ensure_one()
        try:
            import openpyxl
        except ImportError:
            raise UserError(_('The "openpyxl" python library is required to read the AI result file.'))

        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
        Line = self.env['vendor.pricelist.update.line']

        # --- IMPORT ODOO -> Product Matching tab (matched=True) -------------
        import_ws = self._find_sheet(wb, 'import odoo') or wb.active
        import_rows, get = self._sheet_rows_and_getter(import_ws)

        for row in import_rows:
            if not any(row):
                continue

            product_tmpl, display_name = self._resolve_product_template(get(row, 'product_template'))
            uom = self._resolve_uom(get(row, 'unit'), product_tmpl)

            Line.create({
                'update_id': self.id,
                'matched': bool(product_tmpl),
                'vendor_product_code': get(row, 'vendor_product_code'),
                'product_name': display_name or False,
                'product_tmpl_id': product_tmpl.id if product_tmpl else False,
                'price': self._to_float(get(row, 'price')),
                'quantity': self._to_float(get(row, 'min_qty')),
                'uom_id': uom.id if uom else False,
                'discount': self._to_float(get(row, 'discount')),
                'delay': self._to_float(get(row, 'delay')),
                'date_start': self._to_date(get(row, 'date_start')),
                'date_end': self._to_date(get(row, 'date_end')),
                'note': get(row, 'note'),
            })

        # --- UNMATCHED PRODUCTS -> Unmatched Products tab (matched=False) ---
        unmatched_ws = self._find_sheet(wb, 'unmatched products')
        unmatched_rows = []
        if unmatched_ws:
            unmatched_rows, uget = self._sheet_rows_and_getter(unmatched_ws)

            for row in unmatched_rows:
                if not any(row):
                    continue

                raw_name = uget(row, 'product_template')
                product_name = str(raw_name).strip() if raw_name not in (None, '') else False
                uom = self._resolve_uom(uget(row, 'unit'), False)

                Line.create({
                    'update_id': self.id,
                    'matched': False,
                    'vendor_product_code': uget(row, 'vendor_product_code'),
                    'product_name': product_name,
                    'product_tmpl_id': False,
                    'price': self._to_float(uget(row, 'price')),
                    'quantity': self._to_float(uget(row, 'min_qty')),
                    'uom_id': uom.id if uom else False,
                    'discount': self._to_float(uget(row, 'discount')),
                    'delay': self._to_float(uget(row, 'delay')),
                    'date_start': self._to_date(uget(row, 'date_start')),
                    'date_end': self._to_date(uget(row, 'date_end')),
                    'note': uget(row, 'note'),
                })

        if not import_rows and not unmatched_rows:
            raise UserError(_('The AI-generated file is empty.'))

    # =========================================================================
    # STEP 4 - Apply the AI result to the Vendor Pricelist
    # =========================================================================
    def action_apply_to_pricelist(self):
        self.ensure_one()
        if self.state != 'ai_done':
            raise UserError(_('AI analysis must complete successfully before applying changes.'))

        Supplierinfo = self.env['product.supplierinfo']
        created, updated, closed = 0, 0, 0

        matched_lines = self.line_ids.filtered(lambda l: l.matched and l.product_tmpl_id and l.date_start)
        for line in matched_lines:
            domain = [
                ('partner_id', '=', self.vendor_id.id),
                ('product_tmpl_id', '=', line.product_tmpl_id.id),
                ('company_id', '=', self.company_id.id),
            ]

            same_start = Supplierinfo.search(
                domain + [('date_start', '=', line.date_start)], order='create_date desc', limit=1)
            uom = line.uom_id or self._resolve_uom(False, line.product_tmpl_id)

            values = {
                'price': line.price,
                'min_qty': line.quantity or 0.0,
                'uom_id': uom.id if uom else False,
                'discount': line.discount,
                'delay': line.delay,
                'date_end': line.date_end or False,
            }
            values = {k: v for k, v in values.items() if v is not False or k == 'date_end'}

            if same_start and not same_start.date_end:
                same_start.write(values)
                line.write({'apply_state': 'updated', 'applied_supplierinfo_id': same_start.id})
                updated += 1
            else:
                new_line = Supplierinfo.create({
                    'partner_id': self.vendor_id.id,
                    'product_tmpl_id': line.product_tmpl_id.id,
                    'product_id': line.product_id.id if line.product_id else False,
                    'product_code': line.vendor_product_code,
                    'company_id': self.company_id.id,
                    'date_start': line.date_start,
                    'date_end': line.date_end or False,
                    'price': line.price,
                    'min_qty': line.quantity or 0.0,
                    'uom_id': line.uom_id.id if line.uom_id else False,
                    'discount': line.discount,
                    'delay': line.delay,
                })
                line.write({'apply_state': 'created', 'applied_supplierinfo_id': new_line.id})
                created += 1

            other_open = Supplierinfo.search(domain + [
                ('date_start', '!=', line.date_start),
                ('date_end', '=', False),
                ('id', 'not in', (same_start.ids if same_start else []) +
                 (line.applied_supplierinfo_id.ids if line.apply_state == 'created' else [])),
            ], order='create_date desc')
            if other_open:
                previous_day = fields.Date.to_date(line.date_start) - timedelta(days=1)
                other_open.write({'date_end': previous_day})
                closed += len(other_open)

        self.write({'state': 'applied', 'applied_date': fields.Datetime.now()})
        self.message_post(body=_(
            'Applied to Vendor Pricelist: %s updated, %s created, %s previous line(s) closed.'
        ) % (updated, created, closed))
        return True

    # =========================================================================
    # STEP 4b - Resolve Unmatched Products (Map with Existing / Create New)
    # =========================================================================
    def action_process_unmatched(self):
        """For every unmatched line (matched = False) where the user has set
        both an Action (unmatched_action) and a Product (product_id):
        - map_existing: use the selected product's template as-is.
        - create_new: duplicate the selected product's template, named after
          the unmatched line's Product Name, then use that new template.
        In both cases the Vendor Pricelist (product.supplierinfo) line for
        that template/vendor is created or updated exactly like the normal
        Product Matching apply flow (same fields, same "close previous open
        line" behavior), and the source line is marked matched=True."""
        self.ensure_one()

        to_process = self.unmatched_line_ids.filtered(
            lambda l: not l.matched and l.unmatched_action and l.product_id
        )
        if not to_process:
            raise UserError(_(
                'Select an Action and a Product for at least one unmatched line first.'))

        missing_date = to_process.filtered(lambda l: not l.date_start)
        if missing_date:
            raise UserError(_(
                'These lines are missing a Start Date, required to update the Vendor '
                'Pricelist: %s') % ', '.join(missing_date.mapped(lambda l: l.product_name or '')))

        Supplierinfo = self.env['product.supplierinfo']
        created, updated, closed = 0, 0, 0

        for line in to_process:
            if line.unmatched_action == 'create_new':
                new_tmpl = line.product_id.product_tmpl_id.copy({'name': line.product_name})
                product_tmpl = new_tmpl
                product_variant = new_tmpl.product_variant_id
            else:  # map_existing
                product_tmpl = line.product_id.product_tmpl_id
                product_variant = line.product_id

            domain = [
                ('partner_id', '=', self.vendor_id.id),
                ('product_tmpl_id', '=', product_tmpl.id),
                ('company_id', '=', self.company_id.id),
            ]

            same_start = Supplierinfo.search(
                domain + [('date_start', '=', line.date_start)], order='create_date desc', limit=1)
            uom = line.uom_id or self._resolve_uom(False, product_tmpl)

            values = {
                'price': line.price,
                'min_qty': line.quantity or 0.0,
                'uom_id': uom.id if uom else False,
                'discount': line.discount,
                'delay': line.delay,
                'date_end': line.date_end or False,
            }
            values = {k: v for k, v in values.items() if v is not False or k == 'date_end'}

            if same_start and not same_start.date_end:
                same_start.write(values)
                applied = same_start
                apply_state = 'updated'
                updated += 1
            else:
                applied = Supplierinfo.create({
                    'partner_id': self.vendor_id.id,
                    'product_tmpl_id': product_tmpl.id,
                    'product_id': product_variant.id if product_variant else False,
                    'product_code': line.vendor_product_code,
                    'company_id': self.company_id.id,
                    'date_start': line.date_start,
                    'date_end': line.date_end or False,
                    'price': line.price,
                    'min_qty': line.quantity or 0.0,
                    'uom_id': uom.id if uom else False,
                    'discount': line.discount,
                    'delay': line.delay,
                })
                apply_state = 'created'
                created += 1

            other_open = Supplierinfo.search(domain + [
                ('date_start', '!=', line.date_start),
                ('date_end', '=', False),
                ('id', '!=', applied.id),
            ], order='create_date desc')
            if other_open:
                previous_day = fields.Date.to_date(line.date_start) - timedelta(days=1)
                other_open.write({'date_end': previous_day})
                closed += len(other_open)

            line.write({
                'matched': True,
                'product_tmpl_id': product_tmpl.id,
                'product_id': product_variant.id if product_variant else False,
                'apply_state': apply_state,
                'applied_supplierinfo_id': applied.id,
            })

        self.message_post(body=_(
            'Unmatched products resolved: %s updated, %s created, %s previous line(s) closed.'
        ) % (updated, created, closed))
        return True

    def action_apply_all(self):
        """Single button that replaces the separate 'Apply to Vendor Pricelist'
        and 'Map / Create Unmatched Products' buttons. It reuses
        action_apply_to_pricelist() and action_process_unmatched() as-is:
        - Runs action_apply_to_pricelist() once, when the record is still in
          'ai_done' (skipped on a later click once already 'applied', since
          that method itself only allows the 'ai_done' state).
        - Runs action_process_unmatched() whenever there is at least one
          unmatched line with both an Action and a Product set, in either
          state, so it also works for unmatched rows filled in after the
          first apply.
        """
        self.ensure_one()
        if self.state not in ('ai_done', 'applied'):
            raise UserError(_('AI analysis must complete successfully before applying changes.'))

        if self.state == 'ai_done':
            self.action_apply_to_pricelist()

        ready_unmatched = self.unmatched_line_ids.filtered(
            lambda l: not l.matched and l.unmatched_action and l.product_id
        )
        if ready_unmatched:
            self.action_process_unmatched()

        return True


    def action_reset_to_draft(self):
        for rec in self:
            rec.write({
                'state': 'draft',
                'error_message': False,
            })
        return True
